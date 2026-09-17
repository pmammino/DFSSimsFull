"""
build_role_templates.py
=======================
Generate the role-assignment workbooks the playing-time model will read.

    python scripts/build_role_templates.py --target-year 2027

Writes `rosters/player_roles_hitters_<year>.xlsx` and
`rosters/player_roles_pitchers_<year>.xlsx`, pre-populated with every player in
`out/*_pa_projections_<year>.csv` so the sheets are ready to fill in rather
than empty.

Design (see "Designing the role taxonomy" in README_projection_engine.md):

  * **Roles are columns holding probabilities**, not a single label. A player
    in a job battle is 50% full-time / 30% platoon / 20% bench, and one label
    would produce the mean of a bimodal distribution.
  * **Job and timing are separate.** `Role Start` supplies the fraction of the
    season the player holds the job, so "mid-season callup who becomes a
    full-timer" and "...who becomes a platoon bat" are expressible — a single
    "Mid Season Callup" role cannot tell them apart. It also unifies callups
    with return-from-injury: both are "when does this role start".
  * **Availability is orthogonal to role.** A full-time player who misses April
    is Full Time x 0.85 availability, which no role label can express.

The workbooks are live: `Proj PA` / `Proj IP` are SUMPRODUCT formulas over the
role probabilities and the anchors, so editing a probability updates the
projection in the sheet. The pipeline reads the same columns.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from team_context import abbr_for_team_id  # noqa: E402

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
ANCHOR_FILL = PatternFill("solid", fgColor="FFF2CC")   # assumptions
INPUT_FILL = PatternFill("solid", fgColor="FFFFCC")    # cells to fill in
EXAMPLE_FILL = PatternFill("solid", fgColor="E2EFDA")
BLUE = Font(name=FONT, size=10, color="0000FF")        # hardcoded input
BLACK = Font(name=FONT, size=10)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ─────────────────────────────────────────────────────────────────────────────
# Role definitions
# ─────────────────────────────────────────────────────────────────────────────
# PA / IP anchors are per 162 team games, at full availability, with the role
# held all season. They are STARTING POINTS to be fitted from history, not
# measurements — `fit_role_anchors` in the plan is what should replace them.
#
# vL Share is the fraction of plate appearances against left-handed pitching,
# and it is given separately for left- and right-handed batters because the
# same role implies very different exposure depending on which side of a
# platoon the player is on. This is the column that fixes a real bug: the
# pipeline currently derives vL_share from a player's PAST usage
# (splits_model.py:153), so a hitter moving into a platoon job keeps a stale
# league-average share.

HITTER_ROLES: list[dict] = [
    dict(role="Full Time", pa=630, vl_lhb=0.26, vl_rhb=0.29,
         note="Everyday regular, ~150 games. Sees the league handedness mix."),
    dict(role="Everyday DH / 1B-DH", pa=600, vl_lhb=0.26, vl_rhb=0.29,
         note="Rarely rested, no defensive injury exposure."),
    dict(role="Strong Side Platoon", pa=480, vl_lhb=0.12, vl_rhb=0.45,
         note="The high-volume half of a platoon — usually a LHB who sits "
              "against LHP. Plays the majority of games."),
    dict(role="Weak Side Platoon", pa=230, vl_lhb=0.08, vl_rhb=0.68,
         note="The low-volume half — usually a RHB who plays mostly against "
              "LHP. Only ~28% of league PA are vs LHP, which caps his volume."),
    dict(role="Catcher - Primary", pa=500, vl_lhb=0.26, vl_rhb=0.29,
         note="CATCHERS HAVE THEIR OWN LADDER. A full-time catcher is ~500 PA, "
              "not 630 — applying the Full Time anchor over-projects every "
              "catcher in baseball by ~120 PA."),
    dict(role="Catcher - Tandem", pa=340, vl_lhb=0.22, vl_rhb=0.38,
         note="Roughly even split with a partner; often matchup-managed."),
    dict(role="Catcher - Backup", pa=180, vl_lhb=0.24, vl_rhb=0.33,
         note="Clear #2, starts ~45-55 games."),
    dict(role="Utility IF", pa=290, vl_lhb=0.26, vl_rhb=0.31,
         note="Covers multiple infield spots; volume depends on incumbents' "
              "health."),
    dict(role="Utility OF / 4th OF", pa=300, vl_lhb=0.24, vl_rhb=0.33,
         note="Distinct from Utility IF — different volume and platoon usage."),
    dict(role="Bench Bat", pa=150, vl_lhb=0.22, vl_rhb=0.36,
         note="Pinch-hits and spot starts."),
    dict(role="Injury Replacement / 26th Man", pa=90, vl_lhb=0.26, vl_rhb=0.29,
         note="Plays only when someone ahead of him is hurt. His volume is "
              "conditional on OTHER players' injuries — a team-level coupling "
              "that ultimately needs simulating, not a point estimate."),
    dict(role="Depth (no MLB PA)", pa=1, vl_lhb=0.26, vl_rhb=0.29,
         note="Organizational depth. Hits the 1 PA floor (PT_FLOOR_PA) so he "
              "stays present and joinable without moving any aggregate."),
]

PITCHER_ROLES: list[dict] = [
    dict(role="Ace (SP1)", ip=195, gs=32, g=32, sv=0.00, hld=0.00,
         note="Front-line starter, ~6.1 IP per start."),
    dict(role="Mid-Rotation Starter (SP2-3)", ip=170, gs=30, g=30,
         sv=0.00, hld=0.00, note="~5.2 IP per start."),
    dict(role="End-of-Rotation Starter (SP4-5)", ip=130, gs=25, g=26,
         sv=0.00, hld=0.01,
         note="Shorter leash; occasional bullpen appearance."),
    dict(role="Innings-Limited Starter", ip=110, gs=22, g=22,
         sv=0.00, hld=0.00,
         note="Young arm on a workload cap, or a post-surgery ramp. Common "
              "now and materially different from End-of-Rotation."),
    dict(role="Swing Arm / Long Relief", ip=95, gs=8, g=30,
         sv=0.01, hld=0.04, note="Moves between rotation and bullpen."),
    dict(role="Opener", ip=55, gs=18, g=45, sv=0.00, hld=0.03,
         note="Already modelled on the daily side (OPENER_BF_MEAN = 4.6 in "
              "slate_config), so it belongs here for consistency. Takes the "
              "start but faces ~4-5 batters."),
    dict(role="Closer", ip=62, gs=0, g=60, sv=0.65, hld=0.05,
         note="Takes the large majority of the team's save pool. Note IP is "
              "NEARLY FLAT across bullpen roles — what differs is save and "
              "hold context, which is why these roles are worth separating."),
    dict(role="Late Inning RP (Setup)", ip=65, gs=0, g=65, sv=0.12, hld=0.28,
         note="Primary hold earner; fills in for the closer."),
    dict(role="Middle Relief", ip=60, gs=0, g=58, sv=0.04, hld=0.14,
         note="Bridge innings, lower leverage."),
    dict(role="Bullpen Depth Arm", ip=40, gs=0, g=35, sv=0.01, hld=0.05,
         note="Up and down from AAA; mop-up and spot duty."),
    dict(role="Rehab / Injury Return", ip=60, gs=10, g=14, sv=0.02, hld=0.03,
         note="Expected back mid-season. Pair with a Role Start of Mid or "
              "Late rather than discounting the anchor twice."),
    dict(role="Depth (no MLB IP)", ip=1, gs=0, g=1, sv=0.00, hld=0.00,
         note="Organizational depth. Hits the 1 IP floor (PT_FLOOR_IP)."),
]

# Role Start -> share of the season the role is held. This replaces the
# Early/Mid/Late "callup" roles with a strictly more expressive parameter, and
# unifies callup timing with return-from-injury timing — they are the same
# quantity.
TIMING = [
    ("Opening Day", 1.00, "On the roster from day one."),
    ("Early Season (~May)", 0.80, "Called up or activated in April/May."),
    ("Mid Season (~July)", 0.50, "Mid-season callup, or back from a long IL stay."),
    ("Late Season (~Sept)", 0.20, "September callup / late activation."),
]


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 46


def _legend(wb, role_defs, kind):
    ws = wb.create_sheet("Legend", 0)
    ws.sheet_view.showGridLines = False
    vol = "PA" if kind == "hitter" else "IP"
    lines = [
        (f"Role assignments — {kind}s", True),
        ("", False),
        ("WHAT TO EDIT — the yellow cells on the Assignments sheet:", True),
        (f"  * One probability per role column. They should sum to 1.00 for "
         f"each player; the 'Role Prob Sum' column checks this.", False),
        ("  * Role Start — when the player takes the job (see the Timing "
         "table on the Roles sheet).", False),
        ("  * Availability — share of the season he is healthy and on an MLB "
         "roster. 1.00 = full year.", False),
        ("  * Pos / Bats (or Throws) — needed for the typed-slot allocation "
         "and the platoon splits.", False),
        ("", False),
        ("WHY PROBABILITIES AND NOT ONE ROLE:", True),
        ("  A player in a job battle is ~50% full-time / 30% platoon / 20% "
         "bench. Picking one label", False),
        (f"  yields a plausible-looking {vol} figure that is the mean of a "
         f"bimodal distribution — wrong in", False),
        ("  both worlds. Probabilities also give the model the variance, "
         "which is what season-long", False),
        ("  ranking and DFS leverage both want.", False),
        ("", False),
        ("WHY ROLE START IS SEPARATE FROM THE ROLE:", True),
        ("  'Mid-season callup' conflates two things. A callup who becomes a "
         "full-timer and one who", False),
        ("  becomes a platoon bat have very different rate lines. Role x "
         "timing expresses both; one", False),
        ("  combined label cannot. It also means a mid-season callup and a "
         "player back from the IL in", False),
        ("  July are the same parameter.", False),
        ("", False),
        ("WHY AVAILABILITY IS SEPARATE TOO:", True),
        ("  'Injured' is not a role. A full-time player who misses April is "
         "Full Time x 0.85, which no", False),
        ("  role label can say. A season-ending injury gives availability ~0, "
         f"so the role collapses to", False),
        (f"  the 1 {vol} floor — the existing rule, reached through the model "
         f"rather than bolted on.", False),
        ("", False),
        ("COLOUR KEY:", True),
        ("  Yellow = fill this in.   Blue text = a hardcoded anchor you may "
         "tune.   Black = formula.", False),
        ("  Green row = the example; delete it before use.", False),
        ("", False),
        ("FREE AGENTS AND UNDER-PROJECTED TEAMS:", True),
        ("  Set Team to 'FA' for an unsigned player. He keeps a full rate line "
         "and can be given a role", False),
        ("  and playing time, and is excluded from every team aggregate. Then "
         "use the Teams sheet to", False),
        ("  reserve part of a club's budget for the signing you expect it to "
         "make, so that club is", False),
        ("  deliberately UNDER-projected instead of spreading those "
         f"{vol} across the players on hand.", False),
        ("", False),
        ("CALIBRATION STATUS:", True),
        (f"  Every anchor on the Roles sheet is a documented starting point, "
         f"not a fitted value. Fit them", False),
        ("  from history (group real player-seasons by observed role) before "
         "trusting the levels. The", False),
        ("  cross-role ORDERING is more reliable than the absolute numbers.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, size=11, bold=bold)
    ws.column_dimensions["A"].width = 105
    return ws


def _roles_sheet(wb, role_defs, kind):
    ws = wb.create_sheet("Roles")
    if kind == "hitter":
        cols = [("Role", 32), ("PA Anchor", 11), ("vL Share (LHB)", 13),
                ("vL Share (RHB)", 13), ("Definition / notes", 78)]
    else:
        cols = [("Role", 32), ("IP Anchor", 11), ("GS Anchor", 11),
                ("G Anchor", 10), ("Save Share Wt", 13), ("Hold Share Wt", 13),
                ("Definition / notes", 78)]
    for j, (name, width) in enumerate(cols, start=1):
        ws.cell(row=1, column=j, value=name)
        ws.column_dimensions[get_column_letter(j)].width = width
    _style_header(ws, 1, len(cols))

    for i, r in enumerate(role_defs, start=2):
        if kind == "hitter":
            vals = [r["role"], r["pa"], r["vl_lhb"], r["vl_rhb"], r["note"]]
        else:
            vals = [r["role"], r["ip"], r["gs"], r["g"], r["sv"], r["hld"],
                    r["note"]]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BLACK if j in (1, len(vals)) else BLUE
            c.border = BOX
            if j == len(vals):
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(v, float):
                c.number_format = "0.00"
        ws.row_dimensions[i].height = 30

    # Timing table
    base = len(role_defs) + 4
    ws.cell(row=base - 1, column=1, value="Role Start -> Role Share").font = \
        Font(name=FONT, size=11, bold=True)
    for j, name in enumerate(["Role Start", "Role Share", "Meaning"], start=1):
        ws.cell(row=base, column=j, value=name)
    _style_header(ws, base, 3)
    for i, (label, share, note) in enumerate(TIMING, start=base + 1):
        ws.cell(row=i, column=1, value=label).font = BLACK
        c = ws.cell(row=i, column=2, value=share)
        c.font = BLUE
        c.number_format = "0.00"
        ws.cell(row=i, column=3, value=note).font = BLACK
        for j in range(1, 4):
            ws.cell(row=i, column=j).border = BOX

    note_row = base + len(TIMING) + 2
    ws.cell(row=note_row, column=1,
            value="Share weights are RELATIVE within a team — the allocator "
                  "normalizes them to that club's save/hold pool from "
                  "team_wins.py, so they need not sum to 1."
            if kind == "pitcher" else
            "Anchors are per 162 team games at full availability with the role "
            "held all season. Multiply by Role Share and Availability.").font = \
        Font(name=FONT, size=10, italic=True)
    return ws, base


def _full_time_reference(df: pd.DataFrame, col: str = "Last_PA") -> float:
    """PA/TBF a full-time player carries IN THIS FILE.

    Cannot be a fixed 650: `Last_PA` is whatever fraction of a season the
    source data covers, and the committed artifacts were built from a PARTIAL
    2026 season (median 142, max 503). Comparing against absolute thresholds
    would label every regular in baseball a bench player — Ohtani's 289
    partial-season PA came out as "Utility IF" before this.

    The 95th percentile is the reference rather than the max, so one outlier
    cannot compress everyone else.
    """
    if col not in df.columns:
        return 650.0
    v = pd.to_numeric(df[col], errors="coerce").dropna()
    p95 = float(v.quantile(0.95)) if len(v) else 0.0
    return p95 if p95 > 0 else 650.0


def _suggest_hitter_role(row, reference: float):
    """A rough starting point, NOT a projection.

    Scaled to a full-season equivalent so it survives a partial-season source
    file. It knows nothing about position, so it never suggests a catcher role
    — fill those in by hand, and remember a full-time catcher is ~500 PA.
    """
    if row.get("pt_tier") == "floor":
        return "Depth (no MLB PA)"
    pa = row.get("evidence_volume")
    if pa is None or pd.isna(pa):
        pa = row.get("Last_PA", 0) or 0
    share = float(pa) / max(reference, 1.0)      # 1.0 = a full-time workload
    if share >= 0.80:
        return "Full Time"
    if share >= 0.55:
        return "Strong Side Platoon"
    if share >= 0.35:
        return "Utility OF / 4th OF"
    if share >= 0.18:
        return "Bench Bat"
    return "Depth (no MLB PA)"


def _suggest_pitcher_role(row, reference: float, staff_rank=None):
    """Starters split by IP per game (rate-based, so partial seasons are fine).

    Relievers split by their RA9 rank within their own staff, because a bullpen
    role is inherently a within-team standing — the best arm on the staff is
    the likeliest closer. This is a starting point: a real depth chart knows
    things RA9 does not (contract, handedness, the manager's preferences).
    """
    if row.get("pt_tier") == "floor":
        return "Depth (no MLB IP)"
    if str(row.get("role", "")).lower() == "starter":
        ip_g = row.get("weighted_IP_per_G", 0) or 0
        if ip_g >= 5.8:
            return "Ace (SP1)"
        if ip_g >= 5.0:
            return "Mid-Rotation Starter (SP2-3)"
        if ip_g >= 3.5:
            return "End-of-Rotation Starter (SP4-5)"
        return "Swing Arm / Long Relief"
    if staff_rank == 1:
        return "Closer"
    if staff_rank in (2, 3):
        return "Late Inning RP (Setup)"
    if staff_rank is not None and staff_rank <= 6:
        return "Middle Relief"
    return "Bullpen Depth Arm"


def _assignments(wb, players, role_defs, kind, roles_ws_name="Roles"):
    ws = wb.create_sheet("Assignments", 1)
    role_names = [r["role"] for r in role_defs]
    vol = "PA" if kind == "hitter" else "IP"
    hand_col = "Bats" if kind == "hitter" else "Throws"

    last_col = "Last PA" if kind == "hitter" else "Last TBF"
    meta = ["PlayerId", "Name", "Team", "Age", hand_col, "Pos",
            last_col, "Career", "Suggested Role", "Role Start", "Role Share",
            "Availability"]
    derived = [f"Proj {vol}", "Role Prob Sum"]
    if kind == "hitter":
        derived += ["Proj vL Share"]
    else:
        derived += ["Proj GS", "Proj G", "Save Wt", "Hold Wt"]
    headers = meta + role_names + derived + ["Notes"]

    for j, h in enumerate(headers, start=1):
        ws.cell(row=1, column=j, value=h)
    _style_header(ws, 1, len(headers))

    r0 = len(meta) + 1                       # first role column
    r1 = len(meta) + len(role_names)         # last role column
    RC0, RC1 = get_column_letter(r0), get_column_letter(r1)

    # Row 2 — anchors, positioned directly above the role columns they scale,
    # so SUMPRODUCT needs no TRANSPOSE (an array function LibreOffice would
    # only partially evaluate in an openpyxl-written file).
    ws.cell(row=2, column=1, value="ANCHORS ->").font = \
        Font(name=FONT, size=10, bold=True, italic=True)
    key = "pa" if kind == "hitter" else "ip"
    for j, r in enumerate(role_defs, start=r0):
        c = ws.cell(row=2, column=j, value=r[key])
        c.font = BLUE
        c.fill = ANCHOR_FILL
        c.border = BOX
    ws.cell(row=2, column=r1 + 1,
            value=f"<- {vol} per 162 G at full availability").font = \
        Font(name=FONT, size=9, italic=True)
    # Hidden anchor rows for the secondary quantities.
    extra_rows = {}
    if kind == "hitter":
        extra_rows = {"vl_lhb": 3, "vl_rhb": 4}
    else:
        extra_rows = {"gs": 3, "g": 4, "sv": 5, "hld": 6}
    for field, row in extra_rows.items():
        ws.cell(row=row, column=1, value=f"anchor:{field}").font = \
            Font(name=FONT, size=9, italic=True, color="808080")
        for j, r in enumerate(role_defs, start=r0):
            c = ws.cell(row=row, column=j, value=r[field])
            c.font = Font(name=FONT, size=9, color="808080")
            c.number_format = "0.00"
        ws.row_dimensions[row].outlineLevel = 1
        ws.row_dimensions[row].hidden = True

    first_data = max(extra_rows.values()) + 1

    def write_row(i, rec, is_example=False):
        vals = [rec.get(h) for h in meta]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BLACK
            c.border = BOX
            if headers[j - 1] in (hand_col, "Pos", "Role Start",
                                  "Availability"):
                c.fill = EXAMPLE_FILL if is_example else INPUT_FILL
            if headers[j - 1] == "Availability":
                c.number_format = "0.00"
        # Role Share is looked up from the timing table, so the two stay in sync.
        tim_first = len(role_defs) + 5
        tim_last = tim_first + len(TIMING) - 1
        ws.cell(row=i, column=11).value = (
            f'=IFERROR(INDEX({roles_ws_name}!$B${tim_first}:$B${tim_last},'
            f'MATCH(J{i},{roles_ws_name}!$A${tim_first}:$A${tim_last},0)),1)')
        ws.cell(row=i, column=11).number_format = "0.00"

        for j, name in enumerate(role_names, start=r0):
            c = ws.cell(row=i, column=j)
            c.value = rec.get("_probs", {}).get(name)
            c.fill = EXAMPLE_FILL if is_example else INPUT_FILL
            c.font = BLACK
            c.border = BOX
            c.number_format = "0.00"

        col = r1
        col += 1
        ws.cell(row=i, column=col,
                value=f"=SUMPRODUCT({RC0}{i}:{RC1}{i},"
                      f"${RC0}$2:${RC1}$2)*$K{i}*$L{i}").number_format = "0.0"
        col += 1
        ws.cell(row=i, column=col,
                value=f"=SUM({RC0}{i}:{RC1}{i})").number_format = "0.00"
        if kind == "hitter":
            col += 1
            lr = extra_rows["vl_lhb"]
            rr = extra_rows["vl_rhb"]
            pa_col = get_column_letter(r1 + 1)      # the Proj PA column
            # Anchor-weighted, so the blended share reflects where the plate
            # appearances actually come from. The denominator reuses Proj PA
            # (backing out Role Share and Availability) rather than repeating
            # a third SUMPRODUCT — three per row over 600+ rows was enough to
            # time LibreOffice out during recalculation.
            ws.cell(row=i, column=col, value=(
                f'=IFERROR(IF(UPPER($E{i})="L",'
                f'SUMPRODUCT({RC0}{i}:{RC1}{i},${RC0}${lr}:${RC1}${lr},'
                f'${RC0}$2:${RC1}$2),'
                f'SUMPRODUCT({RC0}{i}:{RC1}{i},${RC0}${rr}:${RC1}${rr},'
                f'${RC0}$2:${RC1}$2))'
                f'/({pa_col}{i}/($K{i}*$L{i})),"")'
            )).number_format = "0.00"
        else:
            for field, fmt in (("gs", "0.0"), ("g", "0.0"),
                               ("sv", "0.000"), ("hld", "0.000")):
                col += 1
                ar = extra_rows[field]
                scale = "*$K{0}*$L{0}".format(i) if field in ("gs", "g") else ""
                ws.cell(row=i, column=col, value=(
                    f"=SUMPRODUCT({RC0}{i}:{RC1}{i},"
                    f"${RC0}${ar}:${RC1}${ar}){scale}")).number_format = fmt
        col += 1
        c = ws.cell(row=i, column=col, value=rec.get("Notes"))
        c.fill = EXAMPLE_FILL if is_example else INPUT_FILL
        c.font = BLACK

    # Example row, per the workbook convention: realistic values showing the
    # expected format, clearly marked for deletion.
    if kind == "hitter":
        example = {
            "PlayerId": 999999, "Name": "EXAMPLE — delete this row",
            "Team": "FA", "Age": 28, hand_col: "R", "Pos": "2B",
            "Suggested Role": "Full Time",
            "Role Start": "Early Season (~May)", "Availability": 0.85,
            "Notes": "Unsigned; job battle. 60% full-time / 30% strong-side "
                     "platoon / 10% utility. Missed April.",
            "_probs": {"Full Time": 0.60, "Strong Side Platoon": 0.30,
                       "Utility IF": 0.10},
        }
    else:
        example = {
            "PlayerId": 999999, "Name": "EXAMPLE — delete this row",
            "Team": "FA", "Age": 30, hand_col: "R", "Pos": "RP",
            "Suggested Role": "Late Inning RP (Setup)",
            "Role Start": "Opening Day", "Availability": 0.90,
            "Notes": "70% closer / 30% setup — job not settled.",
            "_probs": {"Closer": 0.70, "Late Inning RP (Setup)": 0.30},
        }
    write_row(first_data, example, is_example=True)

    for i, rec in enumerate(players, start=first_data + 1):
        write_row(i, rec)

    widths = {"Name": 26, "Suggested Role": 28, "Role Start": 20, "Notes": 46}
    for j, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(
            h, 15 if j > len(meta) else 11)
    ws.freeze_panes = ws.cell(row=first_data, column=3)
    return ws


def _teams_sheet(wb, kind):
    ws = wb.create_sheet("Teams")
    share = "pa_share" if kind == "hitter" else "ip_share"
    cols = [("Team", 10), (share, 12), ("Expected signing / note", 70)]
    for j, (n, w) in enumerate(cols, start=1):
        ws.cell(row=1, column=j, value=n)
        ws.column_dimensions[get_column_letter(j)].width = w
    _style_header(ws, 1, len(cols))

    from team_context import TEAM_ABBR_BY_ID
    for i, abbr in enumerate(sorted(TEAM_ABBR_BY_ID.values()), start=2):
        ws.cell(row=i, column=1, value=abbr).font = BLACK
        c = ws.cell(row=i, column=2, value=0.0)
        c.font = BLUE
        c.fill = INPUT_FILL
        c.number_format = "0.00"
        ws.cell(row=i, column=3).fill = INPUT_FILL
        for j in range(1, 4):
            ws.cell(row=i, column=j).border = BOX

    n = len(TEAM_ABBR_BY_ID) + 3
    for k, text in enumerate([
        f"Fraction of each club's {'PA' if kind == 'hitter' else 'IP'} budget "
        f"held back for a signing it has not made yet. 0 to 0.5.",
        "Without a reserve the allocator spreads the full budget across the "
        "players currently on hand, so every incumbent absorbs the playing "
        "time that will actually go to the free agent.",
        "A reserve says 'this club's playing time is not all accounted for'. "
        "It does NOT say the club is better than its roster looks — the "
        "quality of an unsigned player is unknowable, so reserves never touch "
        "the rate projections.",
        "Copy these into the \"reserves\" block of "
        "rosters/team_assignments_<year>.json, which is what the pipeline "
        "reads.",
    ]):
        c = ws.cell(row=n + k, column=1, value=text)
        c.font = Font(name=FONT, size=10, italic=True)
    return ws


def _load(target_year: int, out_dir: Path, kind: str) -> list[dict]:
    side = "hitter" if kind == "hitter" else "pitcher"
    path = out_dir / f"{side}_pa_projections_{target_year}.csv"
    if not path.exists():
        print(f"  {path} not found — emitting an empty template")
        return []
    df = pd.read_csv(path)
    team_col = next((c for c in ("Pred_target_team_id", "Pred_home_team_id",
                                 "home_park_team_id") if c in df.columns), None)
    reference = _full_time_reference(df)
    # Bullpen roles are a within-team standing, so rank relievers by RA9 on
    # their own staff rather than league-wide.
    staff_rank = {}
    if kind == "pitcher" and {"role", "RA9"} <= set(df.columns) and team_col:
        rp = df[df["role"].astype(str).str.lower() == "reliever"]
        for _, g in rp.groupby(team_col):
            for n, (idx, _r) in enumerate(
                    g.sort_values("RA9").iterrows(), start=1):
                staff_rank[int(_r["PlayerId"])] = n
    recs = []
    for _, r in df.iterrows():
        abbr = abbr_for_team_id(r[team_col]) if team_col else None
        recs.append({
            "PlayerId": int(r["PlayerId"]),
            "Name": r.get("Name"),
            "Team": abbr or "",
            "Age": (int(r["Age"]) if pd.notna(r.get("Age")) else None),
            "Bats": r.get("BatSide"),
            "Throws": r.get("PitchHand"),
            "Pos": None,
            "Last PA": (float(r["Last_PA"]) if pd.notna(r.get("Last_PA")) else None),
            "Last TBF": (float(r["Last_PA"]) if pd.notna(r.get("Last_PA")) else None),
            "Career": (float(r["Career_PA"]) if pd.notna(r.get("Career_PA")) else None),
            "Suggested Role": (
                _suggest_hitter_role(r, reference) if kind == "hitter"
                else _suggest_pitcher_role(
                    r, reference, staff_rank.get(int(r["PlayerId"])))),
            "Role Start": "Opening Day",
            "Availability": 1.00,
            "Notes": None,
            "_probs": {},
        })
    return recs


def build(kind: str, target_year: int, out_dir: Path, dest: Path) -> Path:
    role_defs = HITTER_ROLES if kind == "hitter" else PITCHER_ROLES
    players = _load(target_year, out_dir, kind)

    wb = Workbook()
    wb.remove(wb.active)
    _legend(wb, role_defs, kind)
    _assignments(wb, players, role_defs, kind)
    _roles_sheet(wb, role_defs, kind)
    _teams_sheet(wb, kind)
    wb._sheets = [wb["Legend"], wb["Assignments"], wb["Roles"], wb["Teams"]]

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    print(f"  wrote {dest} ({len(players)} {kind}s, {len(role_defs)} roles)")
    return dest


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[3])
    ap.add_argument("--target-year", type=int, default=2027)
    ap.add_argument("--out-dir", type=Path, default=ROOT / "out")
    ap.add_argument("--dest-dir", type=Path, default=ROOT / "rosters")
    a = ap.parse_args(argv)

    for kind in ("hitter", "pitcher"):
        build(kind, a.target_year, a.out_dir,
              a.dest_dir / f"player_roles_{kind}s_{a.target_year}.xlsx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
